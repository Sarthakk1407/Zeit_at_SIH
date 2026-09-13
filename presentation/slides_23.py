"""Slides 2 and 3 — the two flowchart slides.

Kept in their own module because they are the only slides with real diagram
geometry: a fork, a feedback loop, a decision node, two parallel paths and a
merge. Everything is drawn from explicit coordinates so PowerPoint and the
preview renderer agree.
"""
import os
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

def _mentor_image():
    """Find the approval screenshot and crop it to the email body.

    The raw capture is a whole browser window; at slide scale the Gmail
    sidebar and toolbar are unreadable noise. Cropping is done here rather
    than by hand so the build stays reproducible from the original file.
    """
    import glob
    from PIL import Image
    for pat in ("mentor-approval.png", "mentor_approval.png",
                "mentor approval.png", "mentor*approval*.png"):
        hits = [h for h in glob.glob(pat) if "crop" not in h]
        if hits:
            src = hits[0]
            break
    else:
        return None
    out = "mentor-approval-crop.png"
    im = Image.open(src)
    if im.width > 1200:                       # full-window capture
        im = im.crop((330, 214, im.width, 678))
    im.save(out)
    return out, im.width / im.height

from deckkit import (LABEL, INK, STEEL, MID, PALE, PALER, RUST, RUSTPL, GREEN, GREENPL,
                     WHITE, FONT, _in, textbox, block, chip, rule, vrule,
                     pointer_label, stat, flow, flow_label, diamond, band,
                     check_fit, fit_widths, widest,
                     BLUE, BLUE_F, GREEN_F, ORANGE, ORANGE_F, PURPLE, PURPLE_F,
                     CRIMSON, CRIMSON_F, TEAL, TEAL_F, dashed, panel, photo_chip)


# =====================================================================
def slide2(s2, PTS2, title_fn):
    title_fn(s2, "ZEIT — Two-Lane Impulse-Adaptive Noise Cancellation", 25)
    L, R = 0.55, 8.92
    LW, RW = 8.05, 3.86

    textbox(s2, L, 1.20, 12.2, 0.22, [(PTS2[0].upper(), 9, True, LABEL)])
    pointer_label(s2, L, 1.46, LW, PTS2[1])

    # ---------------- lane bands
    band(s2, 2.35, 1.72, 6.25, 0.94, "Lane A",
         "hearing protection  ·  closes in under 1 ms  ·  FxLMS on its own module", RUST)
    band(s2, 2.35, 2.94, 6.25, 0.94, "Lane B",
         "transmit to radio  ·  25.3 ms end to end  ·  eleven blocks on the Pi 5", INK)

    # ---------------- microphones (inputs)
    mics = [("Mic 3 · error", 2.00, RUSTPL, RUST),
            ("Mic 2 · reference", 2.61, PALE, STEEL),
            ("Mic 1 · voice", 3.22, PALE, INK)]
    for t, y, f, e in mics:
        chip(s2, L, y, 1.45, 0.48, t, fill=f, line=e, ts=8.5, tag="s2.mic")

    # ---------------- lane contents
    laneA = ["FxLMS filter", "Anti-noise driver", "The ear"]
    laneB = ["Engine · 11 blocks", "Vocoder · MELPe", "Radio link"]
    xs = [2.55, 4.55, 6.55]
    for x, t in zip(xs, laneA):
        chip(s2, x, 2.00, 1.72, 0.48, t, fill=WHITE, line=RUST, ts=8.5,
             tag="s2.A")
    for x, t in zip(xs, laneB):
        f = GREENPL if t.startswith("Radio") else WHITE
        e = GREEN if t.startswith("Radio") else INK
        chip(s2, x, 3.22, 1.72, 0.48, t, fill=f, line=e, ts=8.5, tag="s2.B")

    # ---------------- flow: inputs into the lanes
    flow(s2, [(2.00, 2.24), (2.55, 2.24)], RUST)                    # mic3 → FxLMS
    flow(s2, [(2.00, 3.46), (2.55, 3.46)], INK)                     # mic1 → engine
    flow(s2, [(2.00, 2.85), (2.22, 2.85), (2.22, 2.34), (2.55, 2.34)], STEEL)
    flow(s2, [(2.22, 2.85), (2.22, 3.36), (2.55, 3.36)], STEEL)     # mic2 → both

    # ---------------- flow: along each lane
    for x in (4.27, 6.27):
        flow(s2, [(x, 2.24), (x + 0.28, 2.24)], RUST)
        flow(s2, [(x, 3.46), (x + 0.28, 3.46)], INK)

    # ---------------- the closed loop that makes Lane A adaptive
    flow(s2, [(7.41, 2.48), (7.41, 2.74), (3.41, 2.74), (3.41, 2.48)], RUST)
    flow_label(s2, 3.90, 2.76, 3.10,
               "error signal — what still reaches the ear", RUST, 7.4)

    # ---------------- the look-ahead, now full width under the lanes
    _lookahead_panel(s2, L, 3.94, LW, 0.76)

    # ---------------- requirement → mechanism
    pointer_label(s2, L, 4.84, LW, PTS2[2])
    rows = [("Impulsive noise — gunshots, artillery", CRIMSON,
             "Detect-and-react transient path, trained through the clipping with α-stable augmentation"),
            ("Adaptive", ORANGE,
             "A live kurtosis classifier retunes γ and β every frame — not tuned once, offline"),
            ("High speech intelligibility", GREEN,
             "A dedicated voice guard, with DNSMOS SIG reported separately so over-suppression stays visible"),
            ("Real-time on embedded hardware", BLUE,
             "23.7 K parameters at 39.6 MMAC/s — it runs on a Raspberry Pi 5, not a Jetson")]
    y = 5.06
    for req, hue, mech in rows:
        block(s2, L, y, LW, 0.36, fill=PALER, radius=0.06)
        block(s2, L, y, 0.07, 0.36, fill=hue, shape=MSO_SHAPE.RECTANGLE)
        textbox(s2, L + 0.16, y + 0.055, 2.86, 0.26, [(req, 8.6, True, hue)],
                tag="s2.req")
        textbox(s2, L + 3.10, y + 0.055, LW - 3.26, 0.26,
                [(mech, 8, False, INK)], line_spacing=0.94)
        y += 0.40

    # ---------------- the prototype itself, then what makes it different
    pointer_label(s2, R, 1.46, RW, PTS2[3])
    block(s2, R, 1.70, RW, 1.34, fill=WHITE, line=MID, line_w=1.0, radius=0.06)
    _picture(s2, "hw/hw-bench.png", R + 0.06, 1.74, RW - 0.12, 1.02)
    textbox(s2, R + 0.10, 2.80, RW - 0.20, 0.20,
            [("THE BENCH AS BUILT — headset, interface, Pi 5, radio", 7, True, INK)],
            align=PP_ALIGN.CENTER)

    diffs = [("Two lanes, not one", GREEN,
              "Protection closes in under a millisecond; transmit tolerates 25 ms. Fusing them makes both impossible."),
             ("Detect, do not predict", ORANGE,
              "Published deep ANC buys latency by predicting noise. A gunshot is unpredictable by definition."),
             ("RMS is undefined here", CRIMSON,
              "Heavy-tailed noise has infinite variance. We scale on percentile and fractional lower-order moments."),
             ("Measured, not downloaded", PURPLE,
              "Calibrated gunshot data of our own, holdout split assigned before capture, checksum-frozen.")]
    y = 3.12
    for i, (h, hue, d) in enumerate(diffs, 1):
        sh = block(s2, R, y, RW, 0.92, fill=WHITE, radius=0.06)
        dashed(sh, hue, 1.0)
        block(s2, R + 0.12, y + 0.12, 0.28, 0.28, fill=hue, radius=0.5,
              shape=MSO_SHAPE.OVAL)
        textbox(s2, R + 0.12, y + 0.155, 0.28, 0.22, [(str(i), 9.5, True, WHITE)],
                align=PP_ALIGN.CENTER)
        textbox(s2, R + 0.50, y + 0.11, RW - 0.62, 0.22, [(h, 9.5, True, hue)],
                tag="s2.diff")
        textbox(s2, R + 0.50, y + 0.34, RW - 0.62, 0.52, [(d, 7.6, False, INK)],
                line_spacing=0.94)
        y += 0.96



def _lookahead_panel(slide, x, y, w, h):
    """Crack before blast — the project's one free source of warning.

    Drawn as a timeline because the claim is a timing claim: both wavefronts
    travel at the same speed, but the shockwave is generated at the bullet's
    closest approach, three metres away, not thirty.
    """
    sh = block(slide, x, y, w, h, fill=ORANGE_F, radius=0.06)
    dashed(sh, ORANGE, 1.1)
    block(slide, x + 0.12, y + 0.10, 0.28, 0.28, fill=ORANGE, radius=0.5,
          shape=MSO_SHAPE.OVAL)
    textbox(slide, x + 0.12, y + 0.135, 0.28, 0.22, [("5", 9.5, True, WHITE)],
            align=PP_ALIGN.CENTER)
    textbox(slide, x + 0.50, y + 0.10, 2.40, 0.22,
            [("Crack before blast", 9.5, True, ORANGE)])
    textbox(slide, x + 0.50, y + 0.32, 2.90, 0.34,
            [("the shockwave starts 3 m away, not 30 — physical look-ahead, "
              "not algorithmic", 7, False, INK)], line_spacing=0.92)

    # time axis: 0 to 100 ms across the right of the panel
    ax0, axw, ay = x + 3.70, w - 4.10, y + 0.40
    rule(slide, ax0, ay, axw, color=INK, weight=1.0)

    def tx(ms):
        return ax0 + (ms / 100.0) * axw

    for ms, lab, col in ((0.0, "0 · trigger", INK),
                         (45.2, "45.2 ms · CRACK", ORANGE),
                         (87.5, "87.5 ms · BLAST", CRIMSON)):
        vrule(slide, tx(ms), ay - 0.08, 0.16, color=col, weight=1.5)
        textbox(slide, tx(ms) - 0.62, ay + 0.11, 1.24, 0.16,
                [(lab, 6.8, True, col)], align=PP_ALIGN.CENTER)

    flow(slide, [(tx(45.2), ay - 0.17), (tx(87.5), ay - 0.17)], ORANGE, 1.4,
         head=False, size=0.10)
    textbox(slide, tx(45.2) - 0.40, ay - 0.36, tx(87.5) - tx(45.2) + 0.80, 0.17,
            [("42.3 ms of warning", 7.5, True, ORANGE)], align=PP_ALIGN.CENTER)


def _picture(slide, path, x, y, w, h):
    """Place an image inside a box, preserving its aspect ratio."""
    import os
    if not os.path.exists(path):
        return None
    from PIL import Image
    iw, ih = Image.open(path).size
    ar = iw / float(ih)
    pw, ph = w, w / ar
    if ph > h:
        ph, pw = h, h * ar
    return slide.shapes.add_picture(path, _in(x + (w - pw) / 2.0),
                                    _in(y + (h - ph) / 2.0),
                                    width=_in(pw), height=_in(ph))


# =====================================================================
def slide3(s3, PTS3, title_fn):
    title_fn(s3, "TECHNICAL APPROACH", 27)

    # ---------------- technologies
    pointer_label(s3, 0.55, 1.16, 12.2, PTS3[0])
    tech = [("C++17", "real-time path"), ("ONNX Runtime", "inference on ARM"),
            ("PyTorch", "training, offline"), ("GTCRN", "23.7 K params"),
            ("Aux-IVA", "blind separation"), ("BMRI", "impulsive path"),
            ("PFFFT", "BSD FFT, not GPL"), ("Raspberry Pi 5", "8 GB, the engine"),
            ("UMC202HD", "2-ch capture")]
    # widths follow the content, so the row reads as sized by hand
    tws = fit_widths(widest(tech), 12.20, 0.10, 8.5, pad=0.34)
    x = 0.55
    for (t, sub), w in zip(tech, tws):
        # the three parts the engine actually ships get the heavier treatment
        key = t in ("GTCRN", "Raspberry Pi 5", "BMRI")
        chip(s3, x, 1.36, w, 0.48, t, sub,
             fill=PALE if key else PALER,
             line=STEEL if key else MID, ts=8.5, ss=6.5, tag="s3.tech")
        x += w + 0.10

    pointer_label(s3, 0.55, 2.02, 12.2, PTS3[1])

    # ---------------- hardware chain
    textbox(s3, 0.55, 2.22, 12.2, 0.18,
            [("HARDWARE CHAIN — no development computer in the signal path",
              7.5, True, LABEL)])
    hw = [("3 microphones", "voice · reference · error"), ("UMC202HD", "24-bit, 48 V"),
          ("Raspberry Pi 5", "blocks 1–11"), ("Communication unit", "MELPe 2400 bps"),
          ("Earcups", "return audio")]
    GAPH = 0.26
    hws = fit_widths(widest(hw), 12.20 - GAPH * (len(hw) - 1) + 0.001,
                     0.0, 8.5, pad=0.46)
    x = 0.55
    for i, ((t, sub), w) in enumerate(zip(hw, hws)):
        chip(s3, x, 2.36, w, 0.46, t, sub, fill=WHITE, line=STEEL, ts=8.5,
             ss=6.5, tag="s3.hw")
        if i < len(hw) - 1:
            flow(s3, [(x + w, 2.59), (x + w + GAPH, 2.59)], STEEL, 1.5, size=0.11)
        x += w + GAPH

    # ---------------- Lane B chain: sequence → decision → two paths → merge
    textbox(s3, 0.55, 2.96, 12.2, 0.18,
            [("LANE B SIGNAL CHAIN — one 20 ms frame, 10 ms hop, 25.3 ms end to end",
              7.5, True, LABEL)])

    row1 = [("01", "Audio in"), ("02", "Aux-IVA"), ("03", "Normalise"),
            ("04", "STFT + ERB")]
    x = 0.55
    for n, t in row1:
        _numbered(s3, x, 3.14, 1.48, 0.48, n, t, PALE, None, INK)
        flow(s3, [(x + 1.48, 3.38), (x + 1.66, 3.38)], STEEL, 1.5, size=0.11)
        x += 1.66

    diamond(s3, 7.10, 2.905, 2.40, 0.95, "05  Classifier", "kurtosis → γ, β",
            ts=8.5, ss=6.5)
    textbox(s3, 9.66, 3.20, 3.10, 0.40,
            [("the adaptive decision — the system retunes per frame "
              "rather than being tuned once, offline", 7.6, True, RUST)],
            line_spacing=0.96)

    # branch out of the decision node
    flow(s3, [(8.30, 3.855), (8.30, 4.00), (6.85, 4.00), (6.85, 4.20)], RUST)
    flow(s3, [(8.30, 4.00), (9.95, 4.00), (9.95, 4.20)], INK)
    flow_label(s3, 5.75, 3.79, 2.20, "impulsive frame", RUST, 7.4)
    flow_label(s3, 9.10, 3.79, 2.60, "stationary / non-stationary", INK, 7.4)

    _numbered(s3, 5.65, 4.20, 2.40, 0.50, "06", "BMRI — detect and interpolate",
              RUSTPL, RUST, INK)
    _numbered(s3, 8.75, 4.20, 2.40, 0.50, "07", "GTCRN — learned mask",
              INK, INK, WHITE)

    # merge the two paths back together
    flow(s3, [(6.85, 4.70), (6.85, 4.90), (4.34, 4.90)], RUST, head=False)
    flow(s3, [(9.95, 4.70), (9.95, 4.90), (4.34, 4.90)], INK, head=False)
    flow(s3, [(4.34, 4.90), (4.34, 5.10)], STEEL)
    flow_label(s3, 4.55, 4.94, 2.20, "recombine", STEEL, 7.4)

    row3 = [("08", "Deep filter"), ("09", "Voice guard"), ("10", "Residual"),
            ("11", "ISTFT"), ("→", "Radio")]
    x = 3.60
    for i, (n, t) in enumerate(row3):
        f, e, fg = (GREENPL, GREEN, INK) if n == "→" else (PALE, None, INK)
        _numbered(s3, x, 5.10, 1.48, 0.48, n, t, f, e, fg)
        if i < len(row3) - 1:
            flow(s3, [(x + 1.48, 5.34), (x + 1.66, 5.34)], STEEL, 1.5, size=0.11)
        x += 1.66

    # ---------------- Lane A, drawn beside Lane B rather than mentioned
    # It shares Mic 2 and nothing else: no Pi, no frame, no neural network.
    textbox(s3, 0.55, 4.14, 4.85, 0.18,
            [("LANE A CONTROL LOOP — SEPARATE MODULE, CLOSES IN UNDER 1 ms",
              7.5, True, RUST)])
    chip(s3, 0.55, 4.36, 1.25, 0.40, "Mic 2", "reference", fill=PALE,
         line=RUST, ts=8, ss=6.2, tag="s3.a")
    chip(s3, 2.10, 4.36, 1.25, 0.40, "FxLMS W(z)", "filtered-x", fill=RUSTPL,
         line=RUST, ts=8, ss=6.2, tag="s3.a")
    chip(s3, 3.65, 4.36, 1.25, 0.40, "Driver → ear", "anti-noise", fill=WHITE,
         line=RUST, ts=8, ss=6.2, tag="s3.a")
    flow(s3, [(1.80, 4.56), (2.10, 4.56)], RUST, 1.4, size=0.10)
    flow(s3, [(3.35, 4.56), (3.65, 4.56)], RUST, 1.4, size=0.10)
    # what is left at the eardrum comes back as the error signal
    flow(s3, [(4.275, 4.76), (4.275, 4.96), (1.175, 4.96), (1.175, 5.00)],
         RUST, 1.25, size=0.10)
    chip(s3, 0.55, 5.00, 1.25, 0.36, "Mic 3", "in-ear error", fill=PALE,
         line=RUST, ts=8, ss=6.2, tag="s3.a")
    flow(s3, [(1.80, 5.18), (2.725, 5.18), (2.725, 4.80)], RUST, 1.25, size=0.10)
    flow_label(s3, 1.86, 5.20, 1.60, "error — the residual", RUST, 6.5)

    # ---------------- offline pipeline, with its validation loop
    textbox(s3, 0.55, 5.72, 12.2, 0.18,
            [("OFFLINE PIPELINE — none of this ships; it shapes the weights that do",
              7.5, True, LABEL)])
    off = [("Range capture", "96 kHz · calibrated"), ("Measure", "16 quantities"),
           ("Validate", "against holdout"), ("Mix + augment", "α-stable, clipped"),
           ("Train", "SI-SNR + L1"), ("Prune + quantise", "structured, INT8"),
           ("Export ONNX", "the only boundary"), ("C++ engine", "parity-tested")]
    GAPO = 0.13
    ows = fit_widths(widest(off, 7.5, 6.2), 12.20 - GAPO * (len(off) - 1) + 0.001,
                     0.0, 7.5, pad=0.30)
    x = 0.55
    ocx = []
    for i, ((t, sub), w) in enumerate(zip(off, ows)):
        chip(s3, x, 5.88, w, 0.44, t, sub, fill=PALER, line=MID, ts=7.5,
             ss=6.2, tag="s3.off")
        if i < len(off) - 1:
            flow(s3, [(x + w, 6.10), (x + w + GAPO, 6.10)], MID, 1.2, size=0.09)
        ocx.append(x + w / 2.0)
        x += w + GAPO

    # retune loop: validation failure feeds back to the generator
    flow(s3, [(ocx[2], 6.32), (ocx[2], 6.48), (ocx[1], 6.48), (ocx[1], 6.32)],
         RUST, 1.25)
    flow_label(s3, ocx[1] - 1.55, 6.50, 3.10,
               "mismatch → retune the generator", RUST, 7.0)

    textbox(s3, 6.20, 6.40, 6.55, 0.20,
            [("25.3 ms algorithmic  ·  10 ms compute per hop  ·  under 1 ms Lane A  "
              "·  39.6 MMAC/s", 8, True, INK)], align=PP_ALIGN.RIGHT)


def _numbered(slide, x, y, w, h, num, text, fill, line, fg):
    sh = block(slide, x, y, w, h, fill=fill, line=line, radius=0.10)
    tf = sh.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(0)
    r = p.add_run(); r.text = num
    r.font.size = Pt(6.5); r.font.bold = True; r.font.name = FONT
    r.font.color.rgb = STEEL if fg == INK else WHITE
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = text
    r2.font.size = Pt(8); r2.font.bold = True; r2.font.name = FONT
    r2.font.color.rgb = fg
    check_fit("blk", text, w, 8, True, lines=2)
    return sh
