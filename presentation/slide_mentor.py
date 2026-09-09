"""The mentor-approval slide.

Built on the template's own instruction slide rather than a fresh one: that
slide already carries the footer bar, the auto-numbering slide-number field,
the footer and the SIH logo, so repurposing it keeps every piece of template
bookkeeping intact. Its instruction content is stripped first.
"""
import glob
import os
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN

from deckkit import (INK, STEEL, MID, PALE, PALER, RUST, WHITE, FONT, _in,
                     textbox, block, rule, check_fit)

KEEP = ("Rectangle", "Slide Number", "Footer Placeholder", "Picture")


def mentor_image():
    """Locate the approval screenshot and crop it to the email body.

    The raw capture is a whole browser window; at slide scale the Gmail
    sidebar and toolbar are unreadable noise. Cropping happens here so the
    build stays reproducible from the original file.
    """
    from PIL import Image
    for pat in ("mentor-approval.png", "mentor_approval.png",
                "mentor approval.png", "mentor*approval*.png"):
        hits = [h for h in glob.glob(pat) if "crop" not in h]
        if hits:
            src = hits[0]
            break
    else:
        return None
    out = "mentor-approval-crop.png"
    im = Image.open(src)
    if im.width > 1200:
        im = im.crop((330, 214, im.width, 678))
    im.save(out)
    return out, im.width / im.height


def build(slide, team, delete_shape):
    # strip the template's instruction content, keep its chrome
    for sh in list(slide.shapes):
        if not sh.name.startswith(KEEP):
            delete_shape(sh)

    # title — this slide has no title placeholder, so one is drawn
    tb = textbox(slide, 0.67, 0.22, 9.5, 0.9,
                 [("MENTOR APPROVAL", 27, True, INK)])
    check_fit("mentor.title", "MENTOR APPROVAL", 9.5, 27, True)

    textbox(slide, 0.67, 1.18, 12.1, 0.22,
            [("GUIDANCE ACCEPTED FOR PROBLEM STATEMENT 26052  ·  INTERNAL SMART "
              "INDIA HACKATHON 2026", 9, True, RUST)])

    # identity row
    block(slide, 0.67, 1.52, 12.1, 0.78, fill=PALER, radius=0.05)
    textbox(slide, 0.87, 1.64, 4.6, 0.34, [("Dr. Swati Singal", 19, True, INK)],
            tag="mentor.name")
    textbox(slide, 0.87, 1.99, 5.6, 0.24,
            [("Assistant Professor, Computer Science — SSCSE", 9.5, False, STEEL)],
            tag="mentor.desig")
    textbox(slide, 6.90, 1.66, 5.7, 0.24,
            [("Sharda University, Greater Noida", 11, True, STEEL)],
            align=PP_ALIGN.RIGHT, tag="mentor.univ")
    textbox(slide, 6.90, 1.99, 5.7, 0.24,
            [("Mentor — Team " + team, 9.5, False, MID)],
            align=PP_ALIGN.RIGHT, tag="mentor.team")

    # the approval itself
    got = mentor_image()
    if got:
        path, ratio = got
        w = 9.00
        h = w / ratio
        x = (13.33 - w) / 2
        y = 2.52
        block(slide, x - 0.09, y - 0.09, w + 0.18, h + 0.18, fill=WHITE,
              line=MID, line_w=0.75, radius=0.03)
        slide.shapes.add_picture(path, _in(x), _in(y), width=_in(w),
                                 height=_in(h))
        cap_y = y + h + 0.16
    else:
        x, w, cap_y = 2.17, 9.00, 5.10
        block(slide, x, 2.52, w, 2.40, fill=WHITE, line=MID, line_w=0.75,
              radius=0.03)
        textbox(slide, x + 0.40, 2.90, w - 0.80, 1.40,
                [("“I am pleased to accept your request and look forward to "
                  "guide and support you all. I would like to meet all your group "
                  "members on Monday afternoon.”", 13, False, INK)],
                line_spacing=1.10)

    textbox(slide, x, cap_y, w, 0.24,
            [("Accepted by email in reply to “Request to Mentor Team " + team +
              " for SIH 2026 — PS 26052”", 8.5, False, MID)],
            align=PP_ALIGN.CENTER, tag="mentor.cap")
