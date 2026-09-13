"""Geometric QA for the deck.

LibreOffice is not available here, so slides cannot be rendered for visual
inspection. This checks the defects a render would have caught, directly from
the shape geometry: content outside the safe area, collisions with the template
chrome, overlapping text boxes, and text that cannot fit its container.

    python qa_geometry.py ZEIT_SIH2026_Idea.pptx
"""
import sys
from pptx import Presentation
from deckkit import text_width_in

EMU = 914400
DECK = sys.argv[1] if len(sys.argv) > 1 else 'ZEIT_SIH2026_Idea.pptx'

# Template chrome that content must not collide with
CHROME = [
    ("footer bar", 0.00, 6.95, 13.33, 0.55),
    ("logo",      10.70, 0.00,  2.46, 1.16),
    # the template's per-slide team badge is deleted by brand(), so it is no
    # longer chrome anything can collide with
]
SAFE = (0.30, 0.00, 13.03, 6.95)   # l, t, r, b


def rect(sh):
    return (sh.left / EMU, sh.top / EMU,
            (sh.left + sh.width) / EMU, (sh.top + sh.height) / EMU)


def overlap(a, b):
    ix = min(a[2], b[2]) - max(a[0], b[0])
    iy = min(a[3], b[3]) - max(a[1], b[1])
    return ix, iy


prs = Presentation(DECK)
problems = 0

for n, slide in enumerate(prs.slides, 1):
    shapes = [s for s in slide.shapes if s.left is not None]
    named = []
    for s in shapes:
        r = rect(s)
        txt = s.text_frame.text.strip() if s.has_text_frame else ""
        named.append((s, r, txt))

    # 1. outside the safe area
    for s, r, txt in named:
        if s.is_placeholder or s.name.startswith(("Rectangle", "Picture", "Oval", "Freeform")):
            continue                      # template chrome, not our content
        if txt and (r[0] < SAFE[0] - 0.01 or r[2] > SAFE[2] + 0.01 or r[3] > SAFE[3] + 0.01):
            print("s%d  OUT-OF-SAFE   %-18s L%.2f R%.2f B%.2f  %s"
                  % (n, s.name[:18], r[0], r[2], r[3], txt[:40]))
            problems += 1

    # 2. collision with template chrome
    for s, r, txt in named:
        if not txt or s.name.startswith(("Rectangle", "Picture", "Oval",
                                         "Slide Number", "Footer", "Title",
                                         "Subtitle", "Freeform")):
            continue
        for cname, cl, ct, cw, ch in CHROME:
            c = (cl, ct, cl + cw, ct + ch)
            ix, iy = overlap(r, c)
            if ix > 0.02 and iy > 0.02:
                print("s%d  HITS %-11s %-18s overlap %.2f x %.2f  %s"
                      % (n, cname, s.name[:18], ix, iy, txt[:40]))
                problems += 1

    # 3. text boxes overlapping each other
    tbs = [(s, r, t) for s, r, t in named
           if t and s.shape_type is not None and "TextBox" in s.name]
    for i in range(len(tbs)):
        for j in range(i + 1, len(tbs)):
            ix, iy = overlap(tbs[i][1], tbs[j][1])
            if ix > 0.06 and iy > 0.06:
                print("s%d  TEXT OVERLAP  %.2f x %.2f  '%s' / '%s'"
                      % (n, ix, iy, tbs[i][2][:26], tbs[j][2][:26]))
                problems += 1

    # 4. single-line text wider than its box
    for s, r, txt in named:
        if not s.has_text_frame or not txt:
            continue
        w = r[2] - r[0]
        h = r[3] - r[1]
        for p in s.text_frame.paragraphs:
            line = "".join(run.text for run in p.runs)
            if not line.strip():
                continue
            size = None
            bold = False
            for run in p.runs:
                if run.font.size:
                    size = run.font.size.pt
                    bold = bool(run.font.bold)
                    break
            if size is None:
                continue
            need = text_width_in(line, size, bold)
            capacity_lines = max(1, int(h / (size * 1.25 / 72.0)))
            if need > (w - 0.10) * capacity_lines:
                print("s%d  TEXT TOO WIDE %-16s needs %.2f\" have %.2f\" x%d  %s"
                      % (n, s.name[:16], need, w - 0.10, capacity_lines, line[:44]))
                problems += 1

print("\nslides: %d   problems: %d" % (len(prs.slides._sldIdLst), problems))
