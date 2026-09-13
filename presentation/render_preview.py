"""Render the deck to PNGs for visual QA.

LibreOffice is not installed in this environment, so the usual
soffice --convert-to pdf route is unavailable. This draws the slides directly
from the shape tree with PIL: fills, outlines, arrows, connectors and text with
real Arial metrics and greedy word wrap.

It is an approximation of PowerPoint's renderer, not a substitute for it — but
it is accurate enough to catch overlap, overflow and alignment faults, which is
what visual QA is for.

    python render_preview.py ZEIT_SIH2026_Idea.pptx
"""
import sys
import os
import io
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

DECK = sys.argv[1] if len(sys.argv) > 1 else 'ZEIT_SIH2026_Idea.pptx'
DPI = 110
EMU = 914400
OUT = 'preview'
os.makedirs(OUT, exist_ok=True)

FONTS = {}


def font(size_pt, bold):
    key = (round(size_pt, 1), bold)
    if key not in FONTS:
        path = "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"
        FONTS[key] = ImageFont.truetype(path, max(6, int(round(size_pt * DPI / 72.0))))
    return FONTS[key]


def px(emu):
    return int(round(emu / EMU * DPI))


def rgb_of(fmt, default=None):
    try:
        if fmt.type is None:
            return default
        c = fmt.fore_color if hasattr(fmt, 'fore_color') else fmt
        if c.type is not None and str(c.type).startswith('MSO_THEME'):
            return default
        return tuple(c.rgb) if c.rgb else default
    except Exception:
        return default


def wrap(draw, text, f, maxw):
    words, lines, cur = text.split(), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if draw.textlength(t, font=f) <= maxw or not cur:
            cur = t
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def draw_shape(d, sh, img):
    if sh.left is None:
        return
    x0, y0 = px(sh.left), px(sh.top)
    x1, y1 = px(sh.left + sh.width), px(sh.top + sh.height)
    st = sh.shape_type

    # --- geometry
    if st == MSO_SHAPE_TYPE.PICTURE:
        try:                                   # draw the real image, not a stub
            src = Image.open(io.BytesIO(sh.image.blob)).convert("RGB")
            img.paste(src.resize((max(1, x1 - x0), max(1, y1 - y0)),
                                 Image.LANCZOS), (x0, y0))
        except Exception:
            d.rectangle([x0, y0, x1, y1], fill=(228, 231, 234))
            d.text((x0 + 6, y0 + 6), "[image]", font=font(8, False),
                   fill=(150, 155, 160))
        return
    if st == MSO_SHAPE_TYPE.LINE:
        try:
            c = tuple(sh.line.color.rgb)
        except Exception:
            c = (140, 150, 160)
        d.line([x0, y0, x1, y1], fill=c, width=2)
        return

    fill = None
    try:
        if sh.fill.type is not None and sh.fill.type == 1:
            fill = tuple(sh.fill.fore_color.rgb)
    except Exception:
        pass
    line = None
    try:
        if sh.line.fill.type == 1:
            line = tuple(sh.line.color.rgb)
    except Exception:
        pass

    name = (sh.name or "")
    auto = ""
    try:
        auto = str(sh.auto_shape_type)
    except Exception:
        pass
    if fill or line:
        if "OVAL" in auto or "Oval" in name:
            d.ellipse([x0, y0, x1, y1], fill=fill, outline=line, width=2)
        elif "TRIANGLE" in auto:
            # arrowheads: rotation says which way the apex points
            rot = int(round(sh.rotation)) % 360
            cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
            hw, hh = (x1 - x0) / 2, (y1 - y0) / 2
            apex = {0: (cx, cy - hh), 90: (cx + hw, cy),
                    180: (cx, cy + hh), 270: (cx - hw, cy)}.get(rot, (cx, cy - hh))
            base = {0: [(cx - hw, cy + hh), (cx + hw, cy + hh)],
                    90: [(cx - hw, cy - hh), (cx - hw, cy + hh)],
                    180: [(cx - hw, cy - hh), (cx + hw, cy - hh)],
                    270: [(cx + hw, cy - hh), (cx + hw, cy + hh)]}.get(
                        rot, [(cx - hw, cy + hh), (cx + hw, cy + hh)])
            d.polygon([apex] + base, fill=fill or line, outline=line)
        elif "DIAMOND" in auto:
            cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
            d.polygon([(cx, y0), (x1, cy), (cx, y1), (x0, cy)],
                      fill=fill, outline=line)
        else:
            try:
                d.rounded_rectangle([x0, y0, x1, y1], radius=6, fill=fill,
                                    outline=line, width=2)
            except Exception:
                d.rectangle([x0, y0, x1, y1], fill=fill, outline=line, width=2)

    # --- text
    if not sh.has_text_frame:
        return
    tf = sh.text_frame
    paras = [p for p in tf.paragraphs]
    if not any(p.text.strip() for p in paras):
        return
    padx = 4
    boxw = (x1 - x0) - 2 * padx
    if "DIAMOND" in auto:
        boxw = int((x1 - x0) * 0.46)      # a diamond gives text far less width
    if boxw < 8:
        return

    rendered = []
    for p in paras:
        txt = "".join(r.text for r in p.runs)
        if not txt.strip():
            rendered.append((None, 0, 0, None, None))
            continue
        size, bold, color = 11.0, False, (30, 32, 36)
        for r in p.runs:
            if r.font.size:
                size = r.font.size.pt
            if r.font.bold is not None:
                bold = bool(r.font.bold)
            try:
                if r.font.color and r.font.color.rgb:
                    color = tuple(r.font.color.rgb)
            except Exception:
                pass
            break
        f = font(size, bold)
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        lh = int(size * DPI / 72.0 * 1.22 * ls)
        for ln in wrap(d, txt, f, boxw):
            rendered.append((ln, lh, size, f, (color, p.alignment)))

    total = sum(r[1] for r in rendered if r[0])
    anchor = tf.vertical_anchor
    if anchor == MSO_ANCHOR.MIDDLE:
        cy = y0 + ((y1 - y0) - total) // 2
    elif anchor == MSO_ANCHOR.BOTTOM:
        cy = y1 - total - 3
    else:
        cy = y0 + 2
    for ln, lh, size, f, meta in rendered:
        if ln is None:
            cy += 6
            continue
        color, align = meta
        w = d.textlength(ln, font=f)
        if align == PP_ALIGN.CENTER:
            tx = x0 + ((x1 - x0) - w) / 2
        elif align == PP_ALIGN.RIGHT:
            tx = x1 - padx - w
        else:
            tx = x0 + padx
        d.text((tx, cy), ln, font=f, fill=color)
        cy += lh


prs = Presentation(DECK)
W, H = px(prs.slide_width), px(prs.slide_height)
paths = []
for n, slide in enumerate(prs.slides, 1):
    img = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(img)
    for sh in slide.shapes:
        try:
            draw_shape(d, sh, img)
        except Exception as e:
            print("  warn s%d %s: %s" % (n, sh.name, e))
    d.rectangle([0, 0, W - 1, H - 1], outline=(200, 205, 210))
    p = os.path.join(OUT, "slide-%d.png" % n)
    img.save(p)
    paths.append(os.path.abspath(p))
    print(p)

print("\n".join(paths))
